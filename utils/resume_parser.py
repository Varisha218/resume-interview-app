"""Extract plain text from an uploaded resume file (.pdf, .docx, .txt)."""

import os
import pdfplumber
import docx


def parse_resume(filepath: str) -> str:
    ext = os.path.splitext(filepath)[1].lower()

    if ext == ".pdf":
        return _parse_pdf(filepath)
    elif ext == ".docx":
        return _parse_docx(filepath)
    elif ext == ".txt":
        return _parse_txt(filepath)
    else:
        raise ValueError(f"Unsupported file type: {ext}. Please upload PDF, DOCX, or TXT.")


def _parse_pdf(filepath: str) -> str:
    text_parts = []
    with pdfplumber.open(filepath) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
    return "\n".join(text_parts).strip()


def _parse_docx(filepath: str) -> str:
    document = docx.Document(filepath)
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
    # Also pull text from tables (some resumes use table layouts)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text)
    return "\n".join(paragraphs).strip()


def _parse_txt(filepath: str) -> str:
    with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
        return f.read().strip()
